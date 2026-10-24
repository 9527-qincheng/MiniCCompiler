#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""MiniC -- 把 LL(1) 语法分析过程（.parse.txt）另存成 Excel 表格。

产物与 .parse.txt 一一对应：<源文件>.parse.xlsx。
版式参考《文法及其相关表格(带语义分析）-新.xlsx》的「语法分析过程」工作表：

    左栏  测试代码（带行号，一格一行）
    右栏  步骤 | 符号栈（逐符号占一格，右对齐，# 在最右）| 单词流（逐记号占一格，
               当前待匹配记号标红）| 动作

用法:
    python tools/parse2xlsx.py tests/parser/ll1_411.mc
    python tools/parse2xlsx.py examples/example.mc --out docs/语法分析过程_示例.xlsx

依赖 openpyxl（pip install openpyxl）。
"""

import argparse
import os
import re
import shutil
import subprocess
import sys
import tempfile

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required (pip install openpyxl)")

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_EXE = os.path.join(REPO_ROOT, "build", "minic.exe")

FONT_CN = "等线"
FONT_MONO = "Consolas"
THIN = Side(style="thin", color="FF000000")
BORDER_THIN = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FILL_HEAD = PatternFill("solid", fgColor="FFF2F2F2")
RED = "FFFF0000"

META_STEPS = re.compile(r"分析步骤:\s*(\d+)")
META_RESULT = re.compile(r"结果:\s*(\S+)")
TOKEN_LINE = re.compile(r"^\s*\d+:\d+\s+(\S+)\s+(.*)$")


# --------------------------------------------------------------------------
# 采集：跑一遍编译器拿到 .token.txt / .parse.txt
# --------------------------------------------------------------------------
def is_ascii(text):
    try:
        text.encode("ascii")
        return True
    except UnicodeEncodeError:
        return False


def run_minic(exe, source, workdir):
    """在临时目录里跑一次编译器，返回 (trace_text, lexemes)。"""
    dst = os.path.join(workdir, "probe.mc")
    shutil.copyfile(source, dst)

    run_exe = exe
    if not is_ascii(exe):
        # 与 GUI 同一个坑：exe 自身路径含中文时进程的文件 I/O 会失效
        run_exe = os.path.join(workdir, "minic.exe")
        shutil.copyfile(exe, run_exe)

    proc = subprocess.run(
        [run_exe, dst, "--dump-tokens", "--dump-parse"],
        cwd=workdir, capture_output=True,
    )
    trace_path = os.path.join(workdir, "probe.parse.txt")
    token_path = os.path.join(workdir, "probe.token.txt")
    if not os.path.isfile(trace_path):
        sys.exit("compiler produced no .parse.txt (exit=%s)\n%s"
                 % (proc.returncode, proc.stderr.decode("utf-8", "replace")[:800]))
    with open(trace_path, encoding="utf-8", errors="replace") as fh:
        trace = fh.read()
    lexemes = read_lexemes(token_path)
    return trace, lexemes


def read_lexemes(path):
    """从 .token.txt 取记号原文（去掉结尾的 EOF）。"""
    lexemes = []
    if not os.path.isfile(path):
        return lexemes
    with open(path, encoding="utf-8", errors="replace") as fh:
        for line in fh:
            m = TOKEN_LINE.match(line.rstrip("\r\n"))
            if not m:
                continue
            kind, lex = m.group(1), m.group(2).rstrip()
            if kind == "EOF":
                break
            lexemes.append(lex)
    return lexemes


def parse_trace(text):
    """| 步骤 | 符号栈 | 剩余输入串 | 动作 | -> [(step, stack[], input[], action)]"""
    rows = []
    for line in text.splitlines():
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if len(cells) != 4 or not cells[0].isdigit():
            continue
        stack = cells[1].strip("`").split()
        inp = cells[2].strip("`").split()
        rows.append((int(cells[0]), stack, inp, cells[3]))
    return rows


def render_input(inp, lexemes):
    """把「标识符 / 整型常量」这类记号名换成源码里的真实文本。"""
    if not lexemes:
        return inp
    full = lexemes + ["#"]
    offset = len(full) - len(inp)
    if offset < 0:
        return inp
    out = []
    for j, name in enumerate(inp):
        idx = offset + j
        out.append(full[idx] if 0 <= idx < len(full) else name)
    return out


def disp_width(text):
    return sum(2 if ord(ch) > 127 else 1 for ch in str(text))


# --------------------------------------------------------------------------
# 生成 xlsx
# --------------------------------------------------------------------------
def build_workbook(source, source_lines, rows, lexemes, steps, result, title):
    # 单词流用真实文本；符号栈保留文法符号名（与设计文档一致）
    data = []
    for step, stack, inp, action in rows:
        data.append((step, stack, render_input(inp, lexemes), action))

    stack_w = max(len(d[1]) for d in data)
    input_w = max(len(d[2]) for d in data)

    c_step = 4                      # D
    c_stack0 = c_step + 1           # E
    c_input0 = c_stack0 + stack_w
    c_action = c_input0 + input_w
    last_col = c_action

    wb = Workbook()
    ws = wb.active
    ws.title = "语法分析过程"

    # --- 标题 -------------------------------------------------------------
    ws.merge_cells(start_row=1, start_column=4, end_row=4, end_column=last_col)
    cell = ws.cell(row=1, column=4, value=title)
    cell.font = Font(name=FONT_CN, size=14, bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    for c in range(4, last_col + 1):
        ws.cell(row=4, column=c).border = Border(bottom=THIN)

    # --- 左栏：测试代码 ---------------------------------------------------
    ws.merge_cells("A1:B1")
    ws["A1"] = "测试代码"
    ws["A1"].font = Font(name=FONT_CN, size=11, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A2"] = "行号"
    ws["A2"].font = Font(name=FONT_CN, size=10)
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    for i, line in enumerate(source_lines):
        r = 3 + i
        n = ws.cell(row=r, column=1, value=i + 1)
        n.font = Font(name=FONT_CN, size=9)
        n.alignment = Alignment(horizontal="center", vertical="center")
        c = ws.cell(row=r, column=2, value=line)
        c.font = Font(name=FONT_MONO, size=9)
        c.alignment = Alignment(horizontal="left", vertical="center")

    # --- 右栏：表头 -------------------------------------------------------
    head_row = 5
    heads = [(c_step, c_step, "步骤"),
             (c_stack0, c_input0 - 1, "符号栈"),
             (c_input0, c_action - 1, "单词流"),
             (c_action, c_action, "动作")]
    for c0, c1, text in heads:
        if c1 > c0:
            ws.merge_cells(start_row=head_row, start_column=c0, end_row=head_row, end_column=c1)
        cell = ws.cell(row=head_row, column=c0, value=text)
        cell.font = Font(name=FONT_CN, size=11, bold=True)
        cell.fill = FILL_HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center")
        for c in range(c0, c1 + 1):
            ws.cell(row=head_row, column=c).border = BORDER_THIN
            ws.cell(row=head_row, column=c).fill = FILL_HEAD

    # --- 右栏：数据 -------------------------------------------------------
    for i, (step, stack, inp, action) in enumerate(data):
        r = head_row + 1 + i
        c = ws.cell(row=r, column=c_step, value=step)
        c.font = Font(name=FONT_CN, size=10)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER_THIN

        pad = stack_w - len(stack)
        for j in range(stack_w):
            col = c_stack0 + j
            val = stack[j - pad] if j >= pad else None
            cell = ws.cell(row=r, column=col, value=val)
            cell.font = Font(name=FONT_CN, size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER_THIN

        for j in range(input_w):
            col = c_input0 + j
            val = inp[j] if j < len(inp) else None
            cell = ws.cell(row=r, column=col, value=val)
            if j == 0:
                cell.font = Font(name=FONT_MONO, size=10, bold=True, color=RED)
            else:
                cell.font = Font(name=FONT_MONO, size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER_THIN

        cell = ws.cell(row=r, column=c_action, value=action)
        cell.font = Font(name=FONT_CN, size=10)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = BORDER_THIN

    # --- 列宽 / 视图 ------------------------------------------------------
    ws.column_dimensions["A"].width = 6
    code_w = max([disp_width(l) for l in source_lines] + [18])
    ws.column_dimensions["B"].width = min(64, code_w + 2)
    ws.column_dimensions["C"].width = 2
    ws.column_dimensions[get_column_letter(c_step)].width = 6

    for j in range(stack_w):
        col = get_column_letter(c_stack0 + j)
        w = max(disp_width(d[1][len(d[1]) - stack_w + j]) if len(d[1]) - stack_w + j >= 0 else 0
                for d in data)
        ws.column_dimensions[col].width = max(5, min(14, w + 2))

    for j in range(input_w):
        col = get_column_letter(c_input0 + j)
        w = max(disp_width(d[2][j]) if j < len(d[2]) else 0 for d in data)
        ws.column_dimensions[col].width = max(4, min(12, w + 2))

    ws.column_dimensions[get_column_letter(c_action)].width = 40
    ws.freeze_panes = "E6"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    return wb


# --------------------------------------------------------------------------
def main(argv=None):
    ap = argparse.ArgumentParser(description="MiniC .parse.txt -> .parse.xlsx")
    ap.add_argument("source", help=".mc 源文件（或已有的 .parse.txt）")
    ap.add_argument("--exe", default=DEFAULT_EXE, help="minic.exe 路径")
    ap.add_argument("--out", default=None, help="输出 xlsx 路径（默认与源文件同目录同名）")
    ap.add_argument("--title", default=None, help="表格标题")
    ap.add_argument("--workdir", default=None,
                    help="编译临时目录（默认系统临时目录；受限环境下可指向可写目录）")
    args = ap.parse_args(argv)

    src = os.path.abspath(args.source)
    if not os.path.isfile(src):
        sys.exit("source not found: %s" % src)

    if src.lower().endswith(".parse.txt"):
        base = src[: -len(".parse.txt")]
        with open(src, encoding="utf-8", errors="replace") as fh:
            trace = fh.read()
        mc = base + ".mc"
        lexemes = read_lexemes(base + ".token.txt")
        code_path = mc if os.path.isfile(mc) else None
    else:
        if not os.path.isfile(args.exe):
            sys.exit("minic.exe not found: %s" % args.exe)
        if args.workdir:
            workdir = os.path.abspath(args.workdir)
            os.makedirs(workdir, exist_ok=True)
            cleanup = False
        else:
            workdir = tempfile.mkdtemp(prefix="minic-xlsx-")
            cleanup = True
        try:
            trace, lexemes = run_minic(args.exe, src, workdir)
        finally:
            if cleanup:
                shutil.rmtree(workdir, ignore_errors=True)
        base = re.sub(r"\.mc$", "", src, flags=re.I)
        code_path = src

    rows = parse_trace(trace)
    if not rows:
        sys.exit("no parse steps found in trace")

    source_lines = []
    if code_path:
        with open(code_path, encoding="utf-8-sig", errors="replace") as fh:
            source_lines = fh.read().splitlines()

    steps = META_STEPS.search(trace)
    result = META_RESULT.search(trace)
    name = os.path.basename(code_path or base)
    title = args.title or "语法分析过程——预测分析法（%s · %s 步 · %s）" % (
        name, steps.group(1) if steps else len(rows), result.group(1) if result else "—")

    wb = build_workbook(src, source_lines, rows, lexemes,
                        steps.group(1) if steps else len(rows),
                        result.group(1) if result else "—", title)
    out = args.out or (base + ".parse.xlsx")
    wb.save(out)
    print("wrote: %s (%d steps)" % (out, len(rows)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
